"""
LegalEase Document Formatter
Handles Microsoft Word (.docx) generation, structured HTML preview formatting,
and document section breakdown.
"""

import html
import io
import re
from typing import Dict, List, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from .text_utils import sanitize_text, parse_terms


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell internal padding in twips (1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_page_number_to_footer(footer):
    """Inserts a dynamic Word PAGE field into the footer paragraph."""
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)

    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    p._p.append(fldSimple)

    run2 = p.add_run(" | LegalEase AI Draft")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(128, 128, 128)


def generate_docx(
    document_text: str,
    document_type: str = "Legal Document",
    effective_date: str = "",
    parties: str = "",
    logo_bytes: Optional[bytes] = None,
) -> bytes:
    """
    Generates a professionally formatted DOCX file with legal typography,
    proper margins, headers/footers, and optional logo.
    """
    sanitized = sanitize_text(document_text)
    doc = Document()

    # Set page margins to 1 inch standard
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

        # Header
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        if logo_bytes:
            try:
                logo_stream = io.BytesIO(logo_bytes)
                header_p.add_run().add_picture(logo_stream, width=Inches(1.5))
            except Exception:
                pass

        # Footer
        footer = section.footer
        add_page_number_to_footer(footer)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(11.5)
    normal_font.color.rgb = RGBColor(30, 41, 59)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(14)
    title_run = title_p.add_run(document_type.upper() if document_type else "LEGAL AGREEMENT")
    title_run.bold = True
    title_run.font.name = 'Times New Roman'
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    # Metadata callout if effective date or parties exist
    if effective_date or parties:
        meta_table = doc.add_table(rows=1, cols=1)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = meta_table.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

        # Background shading for callout box
        shd = parse_xml(r'<w:shd %s w:fill="F1F5F9"/>' % nsdecls('w'))
        cell._tc.get_or_add_tcPr().append(shd)

        meta_p = cell.paragraphs[0]
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_p.paragraph_format.space_before = Pt(2)
        meta_p.paragraph_format.space_after = Pt(2)

        meta_parts = []
        if effective_date:
            meta_parts.append(f"Effective Date: {effective_date}")
        if parties:
            meta_parts.append(f"Parties: {parties}")
        meta_run = meta_p.add_run("  •  ".join(meta_parts))
        meta_run.font.name = 'Times New Roman'
        meta_run.font.size = Pt(10)
        meta_run.font.italic = True
        meta_run.font.color.rgb = RGBColor(71, 85, 105)

        # Spacing after table
        space_p = doc.add_paragraph()
        space_p.paragraph_format.space_before = Pt(6)
        space_p.paragraph_format.space_after = Pt(0)

    # Split lines and construct document elements
    lines = sanitized.split("\n")
    in_signature = False

    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue

        # Detect headings
        is_h1 = line_str.startswith("# ") or (line_str.isupper() and len(line_str) < 45 and not line_str.startswith("SECTION") and not line_str.startswith("ARTICLE"))
        is_h2 = (
            line_str.startswith("## ")
            or re.match(r"^(SECTION|ARTICLE|CLAUSE)\s+\d+", line_str, re.IGNORECASE)
            or (line_str.isupper() and len(line_str) < 60)
        )
        is_h3 = line_str.startswith("### ") or line_str.startswith("#### ") or (re.match(r"^\d+\.\s+[A-Z]", line_str) and len(line_str) < 70)

        if any(term in line_str.upper() for term in ["IN WITNESS WHEREOF", "SIGNATURES", "EXECUTED BY"]):
            in_signature = True

        if is_h1:
            clean_text = re.sub(r"^#+\s*", "", line_str)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_text)
            r.bold = True
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13.5)
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif is_h2:
            clean_text = re.sub(r"^#+\s*", "", line_str)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_text)
            r.bold = True
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(30, 41, 59)
        elif is_h3:
            clean_text = re.sub(r"^#+\s*", "", line_str)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(clean_text)
            r.bold = True
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif line_str.startswith("* ") or line_str.startswith("- ") or line_str.startswith("• "):
            clean_bullet = re.sub(r"^[\*\-•]\s*", "", line_str)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(clean_bullet)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(51, 65, 85)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            # Handle inline bolding (e.g. **Heading:** text)
            parts = re.split(r"(\*\*.*?\*\*)", line_str)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    bold_text = part[2:-2]
                    r = p.add_run(bold_text)
                    r.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(30, 41, 59)

    # Save to BytesIO
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    return docx_buffer.getvalue()


def parse_document_sections(text: str) -> Dict[str, str]:
    """
    Parses a generated document text into structural components:
    title, preamble, clauses, signatures, and disclaimer.
    """
    sanitized = sanitize_text(text)
    lines = sanitized.split("\n")

    sections: Dict[str, str] = {
        "title": "Legal Document",
        "preamble": "",
        "clauses": "",
        "signatures": "",
        "disclaimer": "",
    }

    if not lines:
        return sections

    # Extract title
    title_line = lines[0].strip().lstrip("#").strip()
    sections["title"] = title_line if title_line else "Legal Agreement"

    body_lines = lines[1:]
    current_section = "preamble"
    preamble_lines: List[str] = []
    clauses_lines: List[str] = []
    sig_lines: List[str] = []
    disc_lines: List[str] = []

    for line in body_lines:
        line_upper = line.upper().strip()
        if "DISCLAIMER" in line_upper:
            current_section = "disclaimer"
        elif any(k in line_upper for k in ["IN WITNESS WHEREOF", "SIGNATURES", "FOR AND ON BEHALF"]):
            current_section = "signatures"
        elif any(k in line_upper for k in ["SECTION 1", "ARTICLE 1", "1. DEFINITIONS", "TERMS AND CONDITIONS", "## 1."]):
            current_section = "clauses"

        if current_section == "preamble":
            preamble_lines.append(line)
        elif current_section == "clauses":
            clauses_lines.append(line)
        elif current_section == "signatures":
            sig_lines.append(line)
        elif current_section == "disclaimer":
            disc_lines.append(line)

    sections["preamble"] = "\n".join(preamble_lines).strip()
    sections["clauses"] = "\n".join(clauses_lines).strip()
    sections["signatures"] = "\n".join(sig_lines).strip()
    sections["disclaimer"] = "\n".join(disc_lines).strip()

    return sections


def format_html_preview(text: str) -> str:
    """
    Transforms plain or markdown-formatted legal text into a sleek,
    responsive HTML preview card for the Streamlit UI.
    """
    if not text:
        return "<div class='empty-preview'>No document generated yet.</div>"

    sanitized = sanitize_text(text)
    lines = sanitized.split("\n")

    html_parts = [
        "<div class='legal-doc-sheet'>",
        "<div class='doc-watermark'>LEGAL DRAFT</div>",
    ]

    in_list = False
    in_signature = False

    for line in lines:
        raw_line = line.strip()

        if not raw_line:
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            html_parts.append("<div class='doc-spacer'></div>")
            continue

        # Check for signature section
        if any(term in raw_line.upper() for term in ["IN WITNESS WHEREOF", "SIGNATURES", "FOR AND ON BEHALF"]):
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            in_signature = True
            html_parts.append(f"<div class='doc-sig-header'>{html.escape(raw_line.lstrip('#').strip())}</div>")
            continue

        # Disclaimer Section
        if "DISCLAIMER" in raw_line.upper() and len(raw_line) < 40:
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            html_parts.append(f"<div class='doc-disclaimer-title'>{html.escape(raw_line.lstrip('#').strip())}</div>")
            continue

        # Headings
        if raw_line.startswith("# "):
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            content = html.escape(raw_line[2:].strip())
            html_parts.append(f"<h1 class='doc-title'>{content}</h1>")
        elif raw_line.startswith("## ") or re.match(r"^(SECTION|ARTICLE|CLAUSE)\s+\d+", raw_line, re.IGNORECASE):
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            content = html.escape(re.sub(r"^#+\s*", "", raw_line))
            html_parts.append(f"<h2 class='doc-h2'>{content}</h2>")
        elif raw_line.startswith("### ") or (re.match(r"^\d+\.\s+[A-Z]", raw_line) and len(raw_line) < 70):
            if in_list:
                html_parts.append("</ul>")
                in_list = False
            content = html.escape(re.sub(r"^#+\s*", "", raw_line))
            html_parts.append(f"<h3 class='doc-h3'>{content}</h3>")
        elif raw_line.startswith("* ") or raw_line.startswith("- ") or raw_line.startswith("• "):
            if not in_list:
                html_parts.append("<ul class='doc-bullet-list'>")
                in_list = True
            content = html.escape(re.sub(r"^[\*\-•]\s*", "", raw_line))
            # Format bold within bullet
            content = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", content)
            html_parts.append(f"<li>{content}</li>")
        else:
            if in_list:
                html_parts.append("</ul>")
                in_list = False

            content = html.escape(raw_line)
            content = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", content)

            if in_signature:
                if "___" in content or "Date:" in content or "By:" in content or "Name:" in content:
                    html_parts.append(f"<div class='doc-sig-line'>{content}</div>")
                else:
                    html_parts.append(f"<p class='doc-sig-p'>{content}</p>")
            else:
                html_parts.append(f"<p class='doc-p'>{content}</p>")

    if in_list:
        html_parts.append("</ul>")

    html_parts.append("</div>")
    return "\n".join(html_parts)
